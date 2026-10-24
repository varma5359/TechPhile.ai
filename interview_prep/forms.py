from pathlib import PurePath
from io import BytesIO
from zipfile import ZipFile

from django import forms
from .role_profiles import SENIORITY_LEVELS, profile_choices
from PIL import Image, UnidentifiedImageError


class PaymentSubmissionForm(forms.Form):
    """Collect a payment reference and an optional private receipt file."""

    transaction_reference = forms.CharField(
        label='UPI transaction or UTR reference',
        max_length=128,
        min_length=6,
        widget=forms.TextInput(attrs={
            'class': 'form-control',
            'autocomplete': 'off',
            'placeholder': 'Enter the reference shown in your UPI app',
        }),
    )
    receipt_file = forms.FileField(
        label='Receipt screenshot or PDF',
        required=False,
        widget=forms.ClearableFileInput(attrs={
            'class': 'form-control',
            'accept': '.png,.jpg,.jpeg,.pdf',
        }),
    )

    def clean_transaction_reference(self):
        """Normalize the submitted reference before checking uniqueness."""
        return self.cleaned_data['transaction_reference'].strip().upper()

    def clean_receipt_file(self):
        """Accept only small PNG, JPEG, or PDF receipts with matching contents."""
        receipt = self.cleaned_data.get('receipt_file')
        if not receipt:
            return receipt

        if receipt.size > 5 * 1024 * 1024:
            raise forms.ValidationError('Receipt files must be 5 MB or smaller.')

        extension = PurePath(receipt.name).suffix.lower()
        if extension in {'.jpg', '.jpeg', '.png'}:
            try:
                image = Image.open(receipt)
                if image.format not in {'JPEG', 'PNG'}:
                    raise forms.ValidationError('Upload a PNG or JPEG image, or a PDF.')
                if image.width * image.height > 20_000_000:
                    raise forms.ValidationError('The uploaded image dimensions are too large.')
                image.verify()
            except (Image.DecompressionBombError, UnidentifiedImageError, OSError, ValueError):
                raise forms.ValidationError('The uploaded image could not be verified.')
            finally:
                receipt.seek(0)
        elif extension == '.pdf':
            header = receipt.read(5)
            receipt.seek(0)
            if header != b'%PDF-':
                raise forms.ValidationError('The uploaded file is not a valid PDF receipt.')
        else:
            raise forms.ValidationError('Upload a PNG or JPEG image, or a PDF.')

        return receipt


def _extract_preparation_document(upload, label):
    """Extract readable text from a small PDF, DOCX, or TXT upload."""
    if not upload:
        return None
    if upload.size > 5 * 1024 * 1024:
        raise forms.ValidationError(f'{label} files must be 5 MB or smaller.')

    extension = PurePath(upload.name).suffix.lower()
    raw = upload.read()
    upload.seek(0)
    try:
        if extension == '.pdf':
            if not raw.startswith(b'%PDF-'):
                raise forms.ValidationError('This does not appear to be a valid PDF file.')
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise forms.ValidationError('PDF upload support is unavailable. Please paste the text instead.') from exc
            reader = PdfReader(BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise forms.ValidationError('Password-protected PDFs are not supported. Please upload an unlocked file.')
            if len(reader.pages) > 40:
                raise forms.ValidationError(f'{label} PDFs must be 40 pages or fewer.')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        elif extension == '.docx':
            try:
                with ZipFile(BytesIO(raw)) as package:
                    if '[Content_Types].xml' not in package.namelist() or 'word/document.xml' not in package.namelist():
                        raise forms.ValidationError('This does not appear to be a valid Word document.')
                    if sum(member.file_size for member in package.infolist()) > 20 * 1024 * 1024:
                        raise forms.ValidationError('The uncompressed Word document is too large to process.')
                from docx import Document
            except ImportError as exc:
                raise forms.ValidationError('Word upload support is unavailable. Please paste the text instead.') from exc
            document = Document(BytesIO(raw))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = '\n'.join(paragraphs)
        elif extension == '.txt':
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('cp1252')
        else:
            raise forms.ValidationError('Upload a PDF, Word (.docx), or plain text (.txt) file.')
    except forms.ValidationError:
        raise
    except Exception as exc:
        raise forms.ValidationError('We could not read this file. Try another file or paste the text.') from exc

    text = '\n'.join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise forms.ValidationError('No selectable text was found. Scanned PDFs without selectable text are not supported.')
    if len(text) > 25000:
        raise forms.ValidationError(f'Extracted {label.lower()} text must be 25,000 characters or fewer.')
    return text


class PreparationRequestForm(forms.Form):
    """Collect resume documents or text while requiring provider consent."""

    request_reference = forms.UUIDField(widget=forms.HiddenInput())
    resume_draft_id = forms.IntegerField(required=False, widget=forms.HiddenInput())
    question_target = forms.IntegerField(
        label='How many questions would you like?',
        required=False,
        min_value=30,
        max_value=60,
        initial=45,
        help_text='Choose between 30 and 60. Your selection is saved with this draft.',
        widget=forms.NumberInput(attrs={'class': 'form-control', 'min': 30, 'max': 60, 'step': 1}),
    )
    company_name = forms.CharField(
        max_length=120,
        required=False,
        label='Company (optional)',
        widget=forms.TextInput(attrs={'class': 'form-control', 'autocomplete': 'organization'}),
    )
    job_title = forms.CharField(
        max_length=150,
        widget=forms.TextInput(attrs={'class': 'form-control', 'autocomplete': 'organization-title'}),
    )
    role_family = forms.ChoiceField(
        choices=profile_choices(),
        initial='general-software-engineering',
        widget=forms.Select(attrs={'class': 'form-select'}),
    )
    seniority = forms.ChoiceField(
        choices=[(key, value['label']) for key, value in SENIORITY_LEVELS.items()],
        initial='entry',
        widget=forms.Select(attrs={'class': 'form-select'}),
    )
    job_description = forms.CharField(
        required=False,
        label='Or paste job description',
        max_length=25000,
        widget=forms.Textarea(attrs={'class': 'form-control', 'rows': 9}),
    )
    job_description_file = forms.FileField(
        required=False,
        label='Upload job description file',
        help_text='PDF, Word (.docx), or plain text (.txt), up to 5 MB. Scanned PDFs without selectable text are not supported.',
        widget=forms.ClearableFileInput(attrs={'class': 'form-control', 'accept': '.pdf,.docx,.txt'}),
    )
    resume_file = forms.FileField(
        required=False,
        label='Upload resume file',
        help_text='PDF, Word (.docx), or plain text (.txt), up to 5 MB. Scanned PDFs without selectable text are not supported.',
        widget=forms.ClearableFileInput(attrs={
            'class': 'form-control',
            'accept': '.pdf,.docx,.txt',
        }),
    )
    resume_text = forms.CharField(
        required=False,
        label='Or paste resume text',
        help_text='Include work, projects, skills, education, and transferable experience.',
        max_length=25000,
        widget=forms.Textarea(attrs={'class': 'form-control', 'rows': 12}),
    )
    consent_to_ai_provider = forms.BooleanField(
        label='I agree to send the resume text and job description to the selected question-preparation provider for analysis.',
        required=True,
    )

    def clean_job_description_file(self):
        return _extract_preparation_document(self.cleaned_data.get('job_description_file'), 'Job description')

    def clean_resume_file(self):
        return _extract_preparation_document(self.cleaned_data.get('resume_file'), 'Resume')

    def clean(self):
        """Combine uploaded and pasted materials, requiring both a resume and a JD."""
        cleaned = super().clean()
        for file_name, text_name, label in (
            ('job_description_file', 'job_description', 'job description'),
            ('resume_file', 'resume_text', 'resume'),
        ):
            extracted = cleaned.get(file_name)
            pasted = (cleaned.get(text_name) or '').strip()
            if extracted:
                cleaned[text_name] = '\n\n'.join(part for part in (extracted, pasted) if part)
                if len(cleaned[text_name]) > 25000:
                    self.add_error(file_name, f'Combined {label} text must be 25,000 characters or fewer.')
            elif not pasted:
                self.add_error(text_name, f'Upload a {label} file or paste the {label} text.')
        return cleaned

    def clean_question_target(self):
        """Use the platform default for older clients and validate any choice."""
        return self.cleaned_data.get('question_target') or 45
